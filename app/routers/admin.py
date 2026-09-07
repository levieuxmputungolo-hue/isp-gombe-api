import io
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Header
from sqlalchemy.orm import Session
from sqlalchemy import func
from pydantic import BaseModel
from typing import Optional
import openpyxl

from app.db import get_db
from app.models import Student, Result, University
from app.schemas import ImportRow, ImportResult

router = APIRouter(prefix="/api/admin", tags=["Admin"])

ADMIN_PASSWORD = "isp-gombe-2025"


class AdminLogin(BaseModel):
    password: str


class AdminLoginResponse(BaseModel):
    success: bool
    token: str


class AdminStats(BaseModel):
    total_students: int
    total_results: int
    total_ues: int
    promotions: list[str]
    levels: list[str]
    options: list[str]


class StudentListItem(BaseModel):
    id: int
    matricule: str
    full_name: str
    promotion: str
    level: str
    option: str
    section: str
    annee_academique: str
    result_count: int


class ResultListItem(BaseModel):
    id: int
    student_matricule: str
    student_name: str
    code_ue: str
    intitule_ue: str
    bci: str
    categorie: str
    credits: int
    note: float
    session: str


@router.post("/login", response_model=AdminLoginResponse)
def admin_login(req: AdminLogin):
    if req.password != ADMIN_PASSWORD:
        raise HTTPException(status_code=401, detail="Mot de passe incorrect")
    return AdminLoginResponse(success=True, token="admin-isp-gombe-token")


@router.get("/stats", response_model=AdminStats)
def admin_stats(db: Session = Depends(get_db)):
    total_students = db.query(func.count(Student.id)).scalar() or 0
    total_results = db.query(func.count(Result.id)).scalar() or 0
    ues = db.query(Result.code_ue).distinct().all()
    promotions = db.query(Student.promotion).distinct().all()
    levels = db.query(Student.level).distinct().all()
    options = db.query(Student.option).distinct().all()

    return AdminStats(
        total_students=total_students,
        total_results=total_results,
        total_ues=len(ues),
        promotions=sorted([r[0] for r in promotions if r[0]]),
        levels=sorted([r[0] for r in levels if r[0]]),
        options=sorted([r[0] for r in options if r[0]]),
    )


@router.get("/students", response_model=list[StudentListItem])
def admin_list_students(
    promotion: str = "",
    level: str = "",
    search: str = "",
    db: Session = Depends(get_db),
):
    q = db.query(Student)
    if promotion:
        q = q.filter(Student.promotion == promotion)
    if level:
        q = q.filter(Student.level == level)
    if search:
        q = q.filter(
            Student.full_name.ilike(f"%{search}%") | Student.matricule.ilike(f"%{search}%")
        )
    students = q.order_by(Student.matricule).all()

    result = []
    for s in students:
        rc = db.query(func.count(Result.id)).filter(Result.student_id == s.id).scalar() or 0
        result.append(StudentListItem(
            id=s.id,
            matricule=s.matricule,
            full_name=s.full_name,
            promotion=s.promotion,
            level=s.level,
            option=s.option,
            section=s.section,
            annee_academique=s.annee_academique,
            result_count=rc,
        ))
    return result


@router.get("/results", response_model=list[ResultListItem])
def admin_list_results(
    matricule: str = "",
    code_ue: str = "",
    db: Session = Depends(get_db),
):
    q = db.query(Result).join(Student)
    if matricule:
        q = q.filter(Student.matricule == matricule)
    if code_ue:
        q = q.filter(Result.code_ue == code_ue)
    results = q.order_by(Student.matricule, Result.code_ue).limit(500).all()

    return [
        ResultListItem(
            id=r.id,
            student_matricule=r.student.matricule,
            student_name=r.student.full_name,
            code_ue=r.code_ue,
            intitule_ue=r.intitule_ue,
            bci=r.bci,
            categorie=r.categorie,
            credits=r.credits,
            note=r.note,
            session=r.session,
        )
        for r in results
    ]


@router.delete("/students/{student_id}")
def admin_delete_student(student_id: int, db: Session = Depends(get_db)):
    student = db.query(Student).filter(Student.id == student_id).first()
    if not student:
        raise HTTPException(status_code=404, detail="Etudiant non trouve")
    db.delete(student)
    db.commit()
    return {"message": "Etudiant supprime"}


@router.delete("/results/{result_id}")
def admin_delete_result(result_id: int, db: Session = Depends(get_db)):
    result = db.query(Result).filter(Result.id == result_id).first()
    if not result:
        raise HTTPException(status_code=404, detail="Resultat non trouve")
    db.delete(result)
    db.commit()
    return {"message": "Resultat supprime"}


@router.post("/import-file", response_model=ImportResult)
async def admin_import_file(file: UploadFile = File(...), db: Session = Depends(get_db)):
    if not file.filename:
        raise HTTPException(status_code=400, detail="Aucun fichier fourni")

    ext = file.filename.rsplit(".", 1)[-1].lower() if "." in file.filename else ""
    if ext not in ("xlsx", "xls"):
        raise HTTPException(status_code=400, detail="Format non supporte. Utilisez .xlsx")

    content = await file.read()
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)

    rows_to_import = []
    errors = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        if ws.max_row < 2:
            continue

        headers = [str(c.value).strip().lower() if c.value else "" for c in ws[1]]

        header_map = {}
        for i, h in enumerate(headers):
            if "matricule" in h:
                header_map["matricule"] = i
            elif "nom" in h and "complet" in h:
                header_map["nom_complet"] = i
            elif "nom" in h and "complet" not in h:
                header_map.setdefault("nom_complet", i)
            elif "code" in h and "ue" in h:
                header_map["code_ue"] = i
            elif "intitule" in h or "intitulé" in h:
                header_map["intitule_ue"] = i
            elif "bci" in h:
                header_map["bci"] = i
            elif "cat" in h:
                header_map["categorie"] = i
            elif "credit" in h:
                header_map["credits"] = i
            elif "note" in h and "pond" not in h:
                header_map["note"] = i
            elif "session" in h:
                header_map["session"] = i
            elif "année" in h or "annee" in h:
                header_map["annee_academique"] = i
            elif "promotion" in h:
                header_map["promotion"] = i
            elif "option" in h:
                header_map["option"] = i
            elif "section" in h:
                header_map["section"] = i

        required = ["matricule", "nom_complet", "code_ue", "intitule_ue", "bci", "categorie", "credits", "note"]
        missing = [f for f in required if f not in header_map]
        if missing:
            errors.append(f"Feuille '{sheet_name}': colonnes manquantes: {', '.join(missing)}")
            continue

        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            try:
                def get_val(field):
                    idx = header_map.get(field)
                    if idx is None or idx >= len(row):
                        return ""
                    return row[idx] if row[idx] is not None else ""

                matricule = str(get_val("matricule")).strip()
                if not matricule:
                    continue

                note_val = get_val("note")
                credits_val = get_val("credits")
                if not isinstance(note_val, (int, float)):
                    continue
                if not isinstance(credits_val, (int, float)):
                    continue

                rows_to_import.append(ImportRow(
                    matricule=matricule,
                    nom_complet=str(get_val("nom_complet")).strip(),
                    code_ue=str(get_val("code_ue")).strip(),
                    intitule_ue=str(get_val("intitule_ue")).strip(),
                    bci=str(get_val("bci")).strip(),
                    categorie=str(get_val("categorie")).strip().upper(),
                    credits=int(credits_val),
                    note=float(note_val),
                    session=str(get_val("session")).strip() if get_val("session") else "Normal",
                    annee_academique=str(get_val("annee_academique")).strip() if get_val("annee_academique") else "2024-2025",
                    promotion=str(get_val("promotion")).strip() if get_val("promotion") else "",
                    option=str(get_val("option")).strip() if get_val("option") else "",
                    section=str(get_val("section")).strip() if get_val("section") else "",
                ))
            except Exception as e:
                errors.append(f"Feuille '{sheet_name}' ligne {row_idx}: {str(e)}")

    if not rows_to_import:
        return ImportResult(imported=0, skipped=0, errors=errors + ["Aucune donnee valide trouvee"])

    imported = 0
    skipped = 0

    for row in rows_to_import:
        try:
            note_ponderee = row.note * row.credits

            student = db.query(Student).filter(
                Student.matricule == row.matricule,
                Student.promotion == row.promotion,
                Student.annee_academique == row.annee_academique,
            ).first()

            if not student:
                student = Student(
                    matricule=row.matricule,
                    full_name=row.nom_complet,
                    promotion=row.promotion,
                    level=row.promotion,
                    option=row.option,
                    section=row.section,
                    annee_academique=row.annee_academique,
                )
                db.add(student)
                db.flush()

            existing = db.query(Result).filter(
                Result.student_id == student.id,
                Result.code_ue == row.code_ue,
                Result.bci == row.bci,
                Result.session == row.session,
            ).first()

            if existing:
                existing.note = row.note
                existing.note_ponderee = note_ponderee
                existing.categorie = row.categorie
            else:
                result = Result(
                    student_id=student.id,
                    code_ue=row.code_ue,
                    intitule_ue=row.intitule_ue,
                    bci=row.bci,
                    categorie=row.categorie,
                    credits=row.credits,
                    note=row.note,
                    session=row.session,
                    note_ponderee=note_ponderee,
                )
                db.add(result)

            imported += 1
        except Exception as e:
            errors.append(f"{row.matricule}: {str(e)}")
            skipped += 1

    db.commit()
    return ImportResult(imported=imported, skipped=skipped, errors=errors)
